"""从 /tmp/user_file.xlsx 花名册扫出所有 主+兼 多部门人员, 填 Member.extra_memberships.

extra_memberships 存 JSON: [{"department":"...","title":"...","employment":"兼"}, ...]
只包含 兼 那一行 (主 已经在 Member.department/title)
"""
from __future__ import annotations
import json, re, sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from app.db import SessionLocal
from app.models import Member


KEEP_DEPTS = {
    "公安政法事业部",
    "具身智能事业部",
    "安全情报事业部",
    "教育科技事业部",
    "战略部",
    "科技部",
    "研发部",
}


def clean_name(n: str) -> str:
    return re.sub(r"[（(].*?[)）]", "", n or "").strip()


def main() -> int:
    wb = load_workbook("/tmp/user_file.xlsx", read_only=True, data_only=True)
    ws = wb["组织架构花名册"]
    # name → [(dept, fulltime/兼, title)]
    appearances: dict[str, list[tuple[str, str, str]]] = defaultdict(list)
    for row in ws.iter_rows(values_only=True):
        if not row or not row[0] or not isinstance(row[0], (int, float)):
            continue
        name = str(row[2]).strip()
        dept = str(row[4]).strip() if len(row) > 4 and row[4] else ""
        full_part = str(row[5]).strip() if len(row) > 5 and row[5] else ""
        title = re.sub(r"[（(].*?[)）]", "", str(row[6]).strip() if len(row) > 6 and row[6] else "").strip()
        if not name or not dept:
            continue
        appearances[name].append((dept, full_part, title))

    db = SessionLocal()
    updated = 0
    for raw_name, entries in appearances.items():
        cn = clean_name(raw_name)
        m = db.query(Member).filter_by(name=cn).first()
        if not m:
            continue
        # 找 兼 部门 (在 KEEP_DEPTS 里 且 dept != m.department)
        extras = []
        for dept, fp, title in entries:
            if dept not in KEEP_DEPTS:
                continue
            if dept == m.department:
                continue
            extras.append({"department": dept, "title": title or None, "employment": fp or "兼"})
        if extras:
            m.extra_memberships = json.dumps(extras, ensure_ascii=False)
            print(f"  {cn}: 主={m.department}/{m.title}  额外={extras}")
            updated += 1
        elif m.extra_memberships:
            m.extra_memberships = None
    db.commit()
    print(f"\nupdated {updated} members")
    db.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
